from dataclasses import dataclass
from datetime import date, datetime
from typing import List, Optional, Tuple
from database.connection import get_session
from database.models import Configuracion, CierreCaja, Prestamo, Pago
from services.financial_service import calcular_prestamo, parse_date
from config import CURRENCY_SYMBOL, DATE_FORMAT, DATETIME_FORMAT


@dataclass
class CashSummary:
    capital_inicial: float
    capital_prestado_total: float
    total_cobrado: float
    capital_disponible_caja: float      # (Capital Inicial - Capital Prestado + Total Cobrado)
    capital_en_calle: float             # Capital pendiente de cobro
    saldo_total_pendiente: float        # Saldo total a cobrar (capital + interés + mora)
    ganancia_neta_cobrada: float        # Intereses y moras efectivamente cobradas


def get_capital_inicial() -> float:
    """Retrieves initial capital setting or sets default of 500,000.0."""
    with get_session() as session:
        cfg = session.query(Configuracion).filter_by(clave="capital_inicial").first()
        if cfg:
            try:
                return float(cfg.valor)
            except ValueError:
                return 500000.0
        # Create default
        cfg = Configuracion(clave="capital_inicial", valor="500000.0")
        session.add(cfg)
        return 500000.0


def set_capital_inicial(monto: float) -> bool:
    """Updates initial cash fund configuration."""
    with get_session() as session:
        cfg = session.query(Configuracion).filter_by(clave="capital_inicial").first()
        if cfg:
            cfg.valor = str(round(float(monto), 2))
        else:
            cfg = Configuracion(clave="capital_inicial", valor=str(round(float(monto), 2)))
            session.add(cfg)
    return True


def calcular_balance_caja() -> CashSummary:
    """
    Computes real-time cash balance metrics:
    - Capital líquido disponible en caja = (Capital Inicial - Capital Prestado + Total Cobrado)
    - Capital colocado en la calle = Capital principal pendiente de recuperar
    - Ganancia neta cobrada = Intereses y moras efectivamente cobradas
    """
    capital_ini = get_capital_inicial()

    with get_session() as session:
        prestamos = session.query(Prestamo).all()
        pagos = session.query(Pago).all()

        total_prestado = round(sum(float(p.monto_prestado) for p in prestamos), 2)
        total_cobrado = round(sum(float(pago.monto) for pago in pagos), 2)

        # 1. Capital líquido disponible en caja
        disponible_caja = round(capital_ini - total_prestado + total_cobrado, 2)

        # 2. Capital colocado en la calle y saldos pendientes
        capital_en_calle = 0.0
        saldo_total_pendiente = 0.0
        ganancia_cobrada = 0.0

        for p in prestamos:
            calc = calcular_prestamo(p)
            saldo_total_pendiente += calc.saldo_restante

            # Pagos recibidos para este préstamo
            abonado = calc.total_abonado
            principal = calc.monto_prestado

            # Capital pendiente de amortizar (calle)
            if p.estado == "EN_CURSO":
                pendiente_principal = max(0.0, principal - abonado)
                capital_en_calle += pendiente_principal

            # Ganancia efectivamente cobrada:
            # Los pagos que exceden el capital prestado corresponden a ganancia (interés y mora)
            if abonado > principal:
                ganancia_cobrada += (abonado - principal)

        return CashSummary(
            capital_inicial=capital_ini,
            capital_prestado_total=total_prestado,
            total_cobrado=total_cobrado,
            capital_disponible_caja=disponible_caja,
            capital_en_calle=round(capital_en_calle, 2),
            saldo_total_pendiente=round(saldo_total_pendiente, 2),
            ganancia_neta_cobrada=round(ganancia_cobrada, 2)
        )


def registrar_cierre_diario(fecha_str: str = None, notas: str = "") -> Tuple[bool, str]:
    """
    Freezes and stores in the database the daily cash balance snapshot:
    (saldo inicial, entradas por cobros, salidas por préstamos otorgados y saldo final).
    """
    target_date = fecha_str or date.today().strftime(DATE_FORMAT)

    # Saldo actual en caja
    balance = calcular_balance_caja()
    saldo_final = balance.capital_disponible_caja

    with get_session() as session:
        # Check if already closed today
        existente = session.query(CierreCaja).filter_by(fecha=target_date).first()
        if existente:
            return False, f"Ya existe un cierre registrado para la fecha {target_date} (ID #{existente.id})."

        # Salidas por préstamos otorgados en esa fecha
        prestamos_dia = session.query(Prestamo).filter(Prestamo.fecha_inicio == target_date).all()
        salidas = sum(float(p.monto_prestado) for p in prestamos_dia)

        # Entradas por cobros realizados en esa fecha (buscando por 'YYYY-MM-DD%')
        pagos_dia = session.query(Pago).filter(Pago.fecha_pago.like(f"{target_date}%")).all()
        entradas = sum(float(p.monto) for p in pagos_dia)

        saldo_inicial = round(saldo_final - entradas + salidas, 2)

        cierre = CierreCaja(
            fecha=target_date,
            saldo_inicial=saldo_inicial,
            entradas_cobros=round(entradas, 2),
            salidas_prestamos=round(salidas, 2),
            saldo_final=round(saldo_final, 2),
            notas=notas.strip() or f"Cierre automático diario del {target_date}",
            fecha_registro=datetime.now().strftime(DATETIME_FORMAT)
        )
        session.add(cierre)

    return True, f"Cierre del {target_date} guardado con éxito. Saldo final: {CURRENCY_SYMBOL}{saldo_final:,.2f}"


def obtener_datos_balance_mensual(anio: int, mes: int) -> dict:
    """Aggregates loan and cash activity for a specific month."""
    mes_str = f"{anio}-{mes:02d}"
    balance_actual = calcular_balance_caja()

    with get_session() as session:
        prestamos_mes = session.query(Prestamo).filter(Prestamo.fecha_inicio.like(f"{mes_str}%")).all()
        pagos_mes = session.query(Pago).filter(Pago.fecha_pago.like(f"{mes_str}%")).all()
        cierres_mes = session.query(CierreCaja).filter(CierreCaja.fecha.like(f"{mes_str}%")).order_by(CierreCaja.fecha).all()

        total_prestado = sum(float(p.monto_prestado) for p in prestamos_mes)
        total_cobrado = sum(float(p.monto) for p in pagos_mes)

        # Ganancia neta cobrada en el mes
        ganancia_mes = 0.0
        for pago in pagos_mes:
            # Proporción estimada de interés en el pago
            prestamo = pago.prestamo
            if prestamo:
                calc = calcular_prestamo(prestamo)
                if calc.total_a_pagar > 0:
                    pct_interes = (calc.interes_monto + calc.mora_acumulada) / calc.total_a_pagar
                    ganancia_mes += (pago.monto * pct_interes)

        cierres_data = [
            {
                "fecha": c.fecha,
                "saldo_inicial": c.saldo_inicial,
                "entradas_cobros": c.entradas_cobros,
                "salidas_prestamos": c.salidas_prestamos,
                "saldo_final": c.saldo_final,
                "notas": c.notas or ""
            }
            for c in cierres_mes
        ]

        prestamos_data = [
            {
                "id": p.id,
                "cliente_nombre": p.cliente.nombre_completo if p.cliente else "N/A",
                "fecha_inicio": p.fecha_inicio,
                "fecha_fin": p.fecha_fin,
                "monto_prestado": p.monto_prestado,
                "porcentaje_interes": p.porcentaje_interes,
                "estado": p.estado
            }
            for p in prestamos_mes
        ]

        return {
            "anio": anio,
            "mes": mes,
            "mes_str": mes_str,
            "prestamos_count": len(prestamos_mes),
            "total_prestado": round(total_prestado, 2),
            "pagos_count": len(pagos_mes),
            "total_cobrado": round(total_cobrado, 2),
            "ganancia_estimada_mes": round(ganancia_mes, 2),
            "flujo_neto_mes": round(total_cobrado - total_prestado, 2),
            "capital_disponible_caja": balance_actual.capital_disponible_caja,
            "capital_en_calle": balance_actual.capital_en_calle,
            "cierres": cierres_data,
            "prestamos": prestamos_data
        }


def exportar_balance_mensual_pdf(anio: int, mes: int, file_path: str) -> Tuple[bool, str]:
    """Generates a professional PDF Monthly Balance report using PySide6 QPdfWriter."""
    try:
        from PySide6.QtGui import QPdfWriter, QTextDocument, QPageSize
        from PySide6.QtWidgets import QApplication

        data = obtener_datos_balance_mensual(anio, mes)
        meses_nombres = ["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio",
                         "Julio", "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre"]
        nombre_mes = meses_nombres[mes - 1]

        # Build clean HTML content for PDF
        html = f"""
        <html>
        <head>
            <style>
                body {{ font-family: 'Segoe UI', Arial, sans-serif; color: #1E293B; margin: 20px; }}
                h1 {{ color: #0F172A; font-size: 24px; margin-bottom: 4px; }}
                h2 {{ color: #0284C7; font-size: 16px; margin-top: 20px; margin-bottom: 8px; border-bottom: 1px solid #E2E8F0; padding-bottom: 4px; }}
                .subtitle {{ color: #64748B; font-size: 12px; margin-bottom: 20px; }}
                .kpi-table {{ width: 100%; border-collapse: collapse; margin-bottom: 20px; }}
                .kpi-box {{ background-color: #F8FAFC; border: 1px solid #CBD5E1; border-radius: 6px; padding: 10px; width: 25%; text-align: center; }}
                .kpi-title {{ font-size: 10px; color: #64748B; font-weight: bold; text-transform: uppercase; }}
                .kpi-val {{ font-size: 16px; font-weight: bold; color: #0F172A; margin-top: 4px; }}
                .data-table {{ width: 100%; border-collapse: collapse; margin-top: 10px; font-size: 11px; }}
                .data-table th {{ background-color: #0F172A; color: #FFFFFF; padding: 6px 8px; text-align: left; }}
                .data-table td {{ border-bottom: 1px solid #E2E8F0; padding: 6px 8px; }}
                .total-row {{ font-weight: bold; background-color: #F1F5F9; }}
                .footer {{ margin-top: 30px; font-size: 9px; color: #94A3B8; text-align: right; }}
            </style>
        </head>
        <body>
            <h1>Reporte de Balance Mensual - {nombre_mes} {anio}</h1>
            <div class="subtitle">Consolidado contable de colocación, recupero y arqueos de caja. Generado el {datetime.now().strftime('%d/%m/%Y %H:%M')}.</div>

            <table class="kpi-table">
                <tr>
                    <td class="kpi-box">
                        <div class="kpi-title">Capital Prestado</div>
                        <div class="kpi-val" style="color:#0284C7;">{CURRENCY_SYMBOL}{data['total_prestado']:,.2f}</div>
                    </td>
                    <td class="kpi-box">
                        <div class="kpi-title">Total Cobrado</div>
                        <div class="kpi-val" style="color:#15803D;">{CURRENCY_SYMBOL}{data['total_cobrado']:,.2f}</div>
                    </td>
                    <td class="kpi-box">
                        <div class="kpi-title">Flujo Neto del Mes</div>
                        <div class="kpi-val" style="color:{'#15803D' if data['flujo_neto_mes'] >= 0 else '#B91C1C'};">{CURRENCY_SYMBOL}{data['flujo_neto_mes']:,.2f}</div>
                    </td>
                    <td class="kpi-box">
                        <div class="kpi-title">Ganancia Neta Cobrada</div>
                        <div class="kpi-val" style="color:#16A34A;">{CURRENCY_SYMBOL}{data['ganancia_estimada_mes']:,.2f}</div>
                    </td>
                </tr>
            </table>

            <h2>Historial de Cierres Diarios en el Mes</h2>
            <table class="data-table">
                <thead>
                    <tr>
                        <th>Fecha</th>
                        <th>Saldo Inicial</th>
                        <th>Cobros (+)</th>
                        <th>Préstamos (-)</th>
                        <th>Saldo Final</th>
                        <th>Notas</th>
                    </tr>
                </thead>
                <tbody>
        """

        if data["cierres"]:
            for c in data["cierres"]:
                html += f"""
                <tr>
                    <td>{c['fecha']}</td>
                    <td>{CURRENCY_SYMBOL}{c['saldo_inicial']:,.2f}</td>
                    <td style="color:#15803D;">+{CURRENCY_SYMBOL}{c['entradas_cobros']:,.2f}</td>
                    <td style="color:#0284C7;">-{CURRENCY_SYMBOL}{c['salidas_prestamos']:,.2f}</td>
                    <td><b>{CURRENCY_SYMBOL}{c['saldo_final']:,.2f}</b></td>
                    <td>{c['notas'] or '-'}</td>
                </tr>
                """
        else:
            html += """<tr><td colspan="6" style="text-align:center; color:#64748B; padding:12px;">No se registraron cierres de caja en este mes.</td></tr>"""

        html += f"""
                </tbody>
            </table>

            <h2>Préstamos Otorgados en el Mes ({data['prestamos_count']})</h2>
            <table class="data-table">
                <thead>
                    <tr>
                        <th>ID</th>
                        <th>Cliente</th>
                        <th>Fecha Inicio</th>
                        <th>Vencimiento</th>
                        <th>Monto Prestado</th>
                        <th>Tasa %</th>
                        <th>Estado</th>
                    </tr>
                </thead>
                <tbody>
        """

        for p in data["prestamos"]:
            html += f"""
            <tr>
                <td>#{p['id']}</td>
                <td>{p['cliente_nombre']}</td>
                <td>{p['fecha_inicio']}</td>
                <td>{p['fecha_fin']}</td>
                <td>{CURRENCY_SYMBOL}{p['monto_prestado']:,.2f}</td>
                <td>{p['porcentaje_interes']}%</td>
                <td>{p['estado']}</td>
            </tr>
            """

        html += f"""
                <tr class="total-row">
                    <td colspan="4">TOTAL PRESTADO EN EL MES</td>
                    <td colspan="3">{CURRENCY_SYMBOL}{data['total_prestado']:,.2f}</td>
                </tr>
                </tbody>
            </table>

            <div class="footer">
                Sistema GestionPrestamos v1.0 • Reporte Contable Oficial
            </div>
        </body>
        </html>
        """

        pdf_writer = QPdfWriter(file_path)
        pdf_writer.setPageSize(QPageSize(QPageSize.A4))
        doc = QTextDocument()
        doc.setHtml(html)
        doc.print_(pdf_writer)

        return True, file_path
    except Exception as e:
        return False, f"Error al generar PDF: {str(e)}"


def exportar_balance_mensual_excel(anio: int, mes: int, file_path: str) -> Tuple[bool, str]:
    """Generates monthly balance Excel report with openpyxl."""
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        data = obtener_datos_balance_mensual(anio, mes)
        meses_nombres = ["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio",
                         "Julio", "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre"]
        nombre_mes = meses_nombres[mes - 1]

        wb = openpyxl.Workbook()

        # Sheet 1: Resumen y Cierres
        ws = wb.active
        ws.title = f"Balance {nombre_mes} {anio}"

        header_font = Font(bold=True, color="FFFFFF", name="Segoe UI")
        header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
        title_font = Font(bold=True, size=14, color="0F172A", name="Segoe UI")

        ws.cell(row=1, column=1, value=f"REPORTE MENSUAL: {nombre_mes.upper()} {anio}").font = title_font
        ws.cell(row=2, column=1, value=f"Generado el {datetime.now().strftime('%d/%m/%Y %H:%M')}")

        # KPI Block
        ws.cell(row=4, column=1, value="Capital Prestado Mes").font = Font(bold=True)
        ws.cell(row=4, column=2, value=data["total_prestado"]).number_format = '$#,##0.00'

        ws.cell(row=5, column=1, value="Total Cobrado Mes").font = Font(bold=True)
        ws.cell(row=5, column=2, value=data["total_cobrado"]).number_format = '$#,##0.00'

        ws.cell(row=6, column=1, value="Flujo Neto Mes").font = Font(bold=True)
        ws.cell(row=6, column=2, value=data["flujo_neto_mes"]).number_format = '$#,##0.00'

        ws.cell(row=7, column=1, value="Ganancia Cobrada Estimada").font = Font(bold=True)
        ws.cell(row=7, column=2, value=data["ganancia_estimada_mes"]).number_format = '$#,##0.00'

        # Cierres Diarios
        ws.cell(row=9, column=1, value="CIERRES DIARIOS").font = Font(bold=True, size=12)
        headers = ["Fecha", "Saldo Inicial", "Cobros (+)", "Préstamos (-)", "Saldo Final", "Notas"]
        ws.append([])  # row 10 empty
        ws.append(headers)  # row 11

        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=11, column=c)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")

        r_idx = 12
        for c in data["cierres"]:
            ws.append([
                c["fecha"],
                c["saldo_inicial"],
                c["entradas_cobros"],
                c["salidas_prestamos"],
                c["saldo_final"],
                c["notas"]
            ])
            for col_i in [2, 3, 4, 5]:
                ws.cell(row=r_idx, column=col_i).number_format = '$#,##0.00'
            r_idx += 1

        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

        wb.save(file_path)
        return True, file_path
    except Exception as e:
        return False, f"Error al generar Excel: {str(e)}"
