import csv
import os
import shutil
from datetime import datetime
from pathlib import Path
from typing import List, Tuple
from config import DB_PATH, BACKUPS_DIR, DATETIME_FORMAT
from services.financial_service import calcular_prestamo


def backup_database(destination_dir: str = None) -> Tuple[bool, str]:
    """
    Creates a timestamped snapshot of the sqlite3 database file using shutil.copyfile.
    Returns (success, message_or_filepath).
    """
    if not os.path.exists(DB_PATH):
        return False, "La base de datos aún no existe en disco."

    dest_folder = Path(destination_dir) if destination_dir else BACKUPS_DIR
    dest_folder.mkdir(parents=True, exist_ok=True)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_filename = f"prestamos_backup_{timestamp}.db"
    target_path = dest_folder / backup_filename

    try:
        shutil.copyfile(DB_PATH, target_path)
        return True, str(target_path)
    except Exception as e:
        return False, f"Error al generar backup: {str(e)}"


def export_prestamos_csv(prestamos: list, file_path: str) -> Tuple[bool, str]:
    """Exports loan list with computed financial metrics to CSV."""
    try:
        with open(file_path, mode="w", newline="", encoding="utf-8-sig") as f:
            writer = csv.writer(f)
            writer.writerow([
                "ID Prestamo",
                "Cliente ID",
                "Nombre Completo",
                "Telefono",
                "Estado",
                "Fecha Inicio",
                "Fecha Vencimiento",
                "Tipo Plazo",
                "Duracion",
                "Monto Prestado",
                "Tasa Interes (%)",
                "Total Base",
                "Dias Atraso",
                "Mora Acumulada",
                "Total Abonado",
                "Saldo Restante"
            ])

            for p in prestamos:
                calc = calcular_prestamo(p)
                cliente = p.cliente
                writer.writerow([
                    p.id,
                    p.cliente_id,
                    cliente.nombre_completo if cliente else "N/A",
                    cliente.telefono if cliente else "",
                    p.estado,
                    p.fecha_inicio,
                    p.fecha_fin,
                    p.tipo_plazo,
                    p.duracion_plazo,
                    f"{calc.monto_prestado:.2f}",
                    f"{calc.porcentaje_interes:.2f}",
                    f"{calc.total_base:.2f}",
                    calc.dias_atraso,
                    f"{calc.mora_acumulada:.2f}",
                    f"{calc.total_abonado:.2f}",
                    f"{calc.saldo_restante:.2f}"
                ])
        return True, file_path
    except Exception as e:
        return False, f"Error al exportar CSV: {str(e)}"


def export_prestamos_excel(prestamos: list, file_path: str) -> Tuple[bool, str]:
    """Exports comprehensive loan and payment report to Excel (.xlsx) with styled sheets."""
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        wb = openpyxl.Workbook()

        # Sheet 1: Prestamos
        ws1 = wb.active
        ws1.title = "Préstamos"

        headers1 = [
            "ID", "Cliente", "Teléfono", "Estado", "Inicio", "Vencimiento",
            "Monto Prestado", "Interés %", "Total Base", "Días Mora", "Mora $",
            "Total Pagado", "Saldo Restante"
        ]
        ws1.append(headers1)

        header_font = Font(bold=True, color="FFFFFF", name="Segoe UI")
        header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        thin_border = Border(
            left=Side(style="thin", color="D3D3D3"),
            right=Side(style="thin", color="D3D3D3"),
            top=Side(style="thin", color="D3D3D3"),
            bottom=Side(style="thin", color="D3D3D3")
        )

        for col_num in range(1, len(headers1) + 1):
            cell = ws1.cell(row=1, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        row_idx = 2
        for p in prestamos:
            calc = calcular_prestamo(p)
            cliente = p.cliente
            ws1.append([
                p.id,
                cliente.nombre_completo if cliente else "N/A",
                cliente.telefono if cliente else "",
                p.estado,
                p.fecha_inicio,
                p.fecha_fin,
                calc.monto_prestado,
                calc.porcentaje_interes,
                calc.total_base,
                calc.dias_atraso,
                calc.mora_acumulada,
                calc.total_abonado,
                calc.saldo_restante
            ])

            # Color highlight for overdue or completed
            row_fill = None
            if p.estado == "TERMINADO":
                row_fill = PatternFill(start_color="F2F2F2", fill_type="solid")
            elif calc.dias_atraso > 0:
                row_fill = PatternFill(start_color="FCE4D6", fill_type="solid")  # soft red/orange

            for col_num in range(1, len(headers1) + 1):
                cell = ws1.cell(row=row_idx, column=col_num)
                cell.border = thin_border
                cell.font = Font(name="Segoe UI")
                if row_fill:
                    cell.fill = row_fill

                # Format currency columns
                if col_num in [7, 9, 11, 12, 13]:
                    cell.number_format = '$#,##0.00'
                elif col_num in [8]:
                    cell.number_format = '0.0"%"'

            row_idx += 1

        # Adjust column widths for Sheet 1
        for col in ws1.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws1.column_dimensions[col_letter].width = max(max_len + 3, 12)

        # Sheet 2: Pagos Detallados
        ws2 = wb.create_sheet(title="Historial de Pagos")
        headers2 = ["ID Pago", "ID Préstamo", "Cliente", "Fecha Pago", "Monto", "Método", "Referencia"]
        ws2.append(headers2)

        for col_num in range(1, len(headers2) + 1):
            cell = ws2.cell(row=1, column=col_num)
            cell.font = header_font
            cell.fill = PatternFill(start_color="203764", fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center")

        p_row = 2
        for p in prestamos:
            cliente = p.cliente
            for pago in getattr(p, "pagos", []):
                ws2.append([
                    pago.id,
                    p.id,
                    cliente.nombre_completo if cliente else "N/A",
                    pago.fecha_pago,
                    pago.monto,
                    pago.metodo_pago,
                    pago.comprobante_ref or ""
                ])
                for col_num in range(1, len(headers2) + 1):
                    c = ws2.cell(row=p_row, column=col_num)
                    c.border = thin_border
                    c.font = Font(name="Segoe UI")
                    if col_num == 5:
                        c.number_format = '$#,##0.00'
                p_row += 1

        for col in ws2.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws2.column_dimensions[col_letter].width = max(max_len + 3, 12)

        wb.save(file_path)
        return True, file_path
    except Exception as e:
        return False, f"Error al exportar Excel: {str(e)}"

